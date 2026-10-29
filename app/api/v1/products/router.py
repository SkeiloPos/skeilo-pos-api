from datetime import timedelta
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.db import get_db, get_tenant
from app.core.dependencies import (
    get_current_user,
    require_tenant_admin,
    require_tenant_admin_or_cashier,
)
from app.core.pagination import Page, paginate
from app.core.models import Tenant, User
from app.core.plan_limits import enforce_plan_limit
from app.core.timezone import business_today, utc_now_naive
from app.models.product_import_job import ProductImportJob
from app.api.v1.products.service import ProductService
from app.api.v1.products import import_ as import_service
from app.api.v1.products.export import build_products_excel, build_import_errors_excel
from app.celery_task import import_products_task
from app.api.v1.products.schemas import (
    ProductCreate,
    ProductUpdate,
    ProductAvailabilityUpdate,
    ProductResponse,
    ProductListResponse,
    ProductDetailResponse,
    ProductSaveResponse,
    ProductImportJobResponse,
    PRODUCT_IMPORT_SHEETS,
)

router = APIRouter(prefix="/products", tags=["products"])

service = ProductService()

# research.md D7: mismo criterio que inventory.service.IMPORT_JOB_TIMEOUT -- un
# job "procesando" más viejo que esto se considera abandonado.
IMPORT_JOB_TIMEOUT = timedelta(minutes=30)


def _find_blocking_import_job(db: Session) -> ProductImportJob | None:
    candidatos = db.execute(
        select(ProductImportJob).where(ProductImportJob.status == "procesando")
    ).scalars().all()
    ahora = utc_now_naive()
    for job in candidatos:
        if ahora - job.started_at > IMPORT_JOB_TIMEOUT:
            job.status = "interrumpido"
            job.finished_at = ahora
            db.add(job)
        else:
            db.commit()
            return job
    db.commit()
    return None


def _job_response(job: ProductImportJob) -> ProductImportJobResponse:
    job.has_errors = job.failed_count > 0
    return ProductImportJobResponse.model_validate(job)


def _get_import_job_or_404(db: Session, job_id: UUID) -> ProductImportJob:
    job = db.get(ProductImportJob, job_id)
    if job is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Import job not found")
    return job


@router.get("/export", summary="Exportar la carta a .xlsx")
def export_products(
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_tenant),
    _: User = Depends(require_tenant_admin),
):
    buffer = build_products_excel(db)
    filename = f"carta_{business_today(tenant).isoformat()}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post(
    "/import", response_model=ProductImportJobResponse,
    status_code=status.HTTP_202_ACCEPTED, summary="Importar la carta desde .xlsx",
)
def import_products(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_tenant),
    user: User = Depends(require_tenant_admin),
):
    try:
        wb = load_workbook(file.file, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY, "El archivo no es un .xlsx válido"
        ) from exc

    hojas_headers: dict[str, list] = {}
    hojas_filas: dict[str, list[list]] = {}
    for nombre_hoja in wb.sheetnames:
        ws = wb[nombre_hoja]
        filas_wb = list(ws.iter_rows(values_only=True))
        if not filas_wb:
            continue
        hojas_headers[nombre_hoja] = list(filas_wb[0])
        hojas_filas[nombre_hoja] = [list(f) for f in filas_wb[1:]]

    problemas = import_service.validar_estructura(hojas_headers)
    if problemas:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "; ".join(problemas))

    en_curso = _find_blocking_import_job(db)
    if en_curso is not None:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"Ya hay una importación de la carta en curso (job {en_curso.id})",
        )

    hojas_filas = {k: v for k, v in hojas_filas.items() if k in PRODUCT_IMPORT_SHEETS}
    total_rows = sum(len(filas) for filas in hojas_filas.values())

    job = ProductImportJob(
        status="procesando", file_name=file.filename or "carta.xlsx",
        started_by_user_id=user.id, total_rows=total_rows,
    )
    db.add(job)
    db.commit()
    db.refresh(job)

    import_products_task.delay(tenant.schema, str(job.id), hojas_filas)
    return _job_response(job)


@router.get(
    "/import/current", response_model=ProductImportJobResponse | None,
    summary="Job de importación de la carta más reciente del negocio",
)
def current_product_import_job(
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    job = db.execute(
        select(ProductImportJob).order_by(ProductImportJob.started_at.desc())
    ).scalars().first()
    return _job_response(job) if job is not None else None


@router.get(
    "/import/{job_id}", response_model=ProductImportJobResponse,
    summary="Job de importación de la carta puntual",
)
def get_product_import_job(
    job_id: UUID,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    return _job_response(_get_import_job_or_404(db, job_id))


@router.get(
    "/import/{job_id}/errors", summary="Reporte de filas fallidas de un import de la carta",
)
def download_product_import_errors(
    job_id: UUID,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    job = _get_import_job_or_404(db, job_id)
    if job.failed_count == 0:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Este import no tiene filas fallidas")

    buffer = build_import_errors_excel(job.errors)
    filename = f"errores_{job.file_name}"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "",
    response_model=Page[ProductListResponse],
    summary="Listar productos",
    description="Devuelve los productos de forma paginada (más recientes primero). Filtra por estado activo y/o disponibilidad.",
    responses={401: {"description": "No autenticado o token inválido."}},
)
def list_products(
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    active: bool | None = Query(None, description="Filtra por estado activo/inactivo."),
    available: bool | None = Query(
        None, description="Filtra por disponibilidad (spec 093: Disponibles/Agotados)."
    ),
    search: str | None = Query(None, description="Búsqueda por nombre."),
    is_topping: bool | None = Query(
        None,
        description=(
            "spec 129: filtra por si el producto está marcado como adición "
            "(topping) -- usado por el selector de 'Adiciones disponibles'."
        ),
    ),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    return paginate(db, service.list_query(active, search, available, is_topping), page, size)


@router.get(
    "/{id}",
    response_model=ProductDetailResponse,
    summary="Obtener un producto",
    responses={
        401: {"description": "No autenticado o token inválido."},
        404: {"description": "El producto no existe."},
    },
)
def get_product(
    id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    # Spec 103-detalle-producto-consolidado: `get_detail_or_404` precarga receta/grupos
    # de opciones por variante (research.md D2); `to_detail_response` necesita el `user`
    # para ocultarle la receta al rol Cajero (D3, FR-003).
    return service.to_detail_response(service.get_detail_or_404(db, id), user)


@router.patch(
    "/{id}/availability",
    response_model=ProductResponse,
    summary="Marcar o desmarcar un producto como agotado",
    description=(
        "Único endpoint por el que Cajero y Administrador tocan `available` (spec 093). "
        "A diferencia de `PATCH /products/{id}`, no acepta ningún otro campo."
    ),
    responses={
        401: {"description": "No autenticado o token inválido."},
        403: {"description": "El usuario no es Administrador ni Cajero."},
        404: {"description": "El producto no existe."},
        409: {"description": "El producto está inactivo: no se puede cambiar su disponibilidad."},
    },
)
def set_product_availability(
    id: UUID,
    body: ProductAvailabilityUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_tenant_admin_or_cashier),
):
    return service.set_availability(db, id, body.available, user)


@router.post(
    "",
    response_model=ProductSaveResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Crear un producto",
    description=(
        "Crea un producto de catálogo junto con sus presentaciones, receta y grupos de opciones "
        "en una sola operación atómica (spec 043). Sin `variants`, recibe una variante default."
    ),
    responses={
        401: {"description": "No autenticado o token inválido."},
        404: {"description": "La categoría, un insumo o un grupo de opciones no existen."},
        409: {"description": "Nombre de presentación o SKU ya en uso."},
        422: {"description": "Datos de entrada inválidos."},
    },
)
def create_product(
    body: ProductCreate,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_tenant),
    _: User = Depends(require_tenant_admin),
):
    enforce_plan_limit(db, tenant, "productos")  # spec 033, FR-005/FR-006
    product = service.create_product(db, tenant, body)
    return service.to_save_response(product)


@router.patch(
    "/{id}",
    response_model=ProductSaveResponse,
    summary="Actualizar un producto",
    description=(
        "Actualiza un producto y, opcionalmente, el árbol completo de sus presentaciones "
        "(crear/editar/desactivar, receta, grupos de opciones y orden) en una sola operación "
        "atómica (spec 043)."
    ),
    responses={
        401: {"description": "No autenticado o token inválido."},
        404: {"description": "El producto, la categoría, un insumo o un grupo de opciones no existen."},
        409: {"description": "Nombre de presentación o SKU ya en uso."},
        422: {"description": "Datos de entrada inválidos."},
    },
)
def update_product(
    id: UUID,
    body: ProductUpdate,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_tenant),
    _: User = Depends(require_tenant_admin),
):
    product = service.update_product(db, tenant, id, body)
    return service.to_save_response(product)


@router.put(
    "/{id}",
    response_model=ProductSaveResponse,
    summary="Actualizar un producto (alias de PATCH)",
    responses={
        401: {"description": "No autenticado o token inválido."},
        404: {"description": "El producto, la categoría, un insumo o un grupo de opciones no existen."},
        409: {"description": "Nombre de presentación o SKU ya en uso."},
        422: {"description": "Datos de entrada inválidos."},
    },
)
def replace_product(
    id: UUID,
    body: ProductUpdate,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_tenant),
    _: User = Depends(require_tenant_admin),
):
    product = service.update_product(db, tenant, id, body)
    return service.to_save_response(product)


@router.delete(
    "/{id}",
    response_model=ProductResponse,
    summary="Desactivar un producto",
    description="Borrado lógico: marca el producto como inactivo (active=False).",
    responses={
        401: {"description": "No autenticado o token inválido."},
        404: {"description": "El producto no existe."},
    },
)
def delete_product(
    id: UUID,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    return service.soft_delete(db, id)
