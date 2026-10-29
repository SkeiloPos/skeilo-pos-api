"""Generación del `.xlsx` de la carta (spec 143, research.md D8): 4 hojas
relacionadas (Productos, Presentaciones y precios, Receta, Adiciones) en un solo
workbook, armado completo en memoria antes de responder -- mismo criterio que
`inventory/export.py` (spec 086/109): un error a mitad de la generación nunca deja
un archivo corrupto/incompleto en la respuesta."""
import io

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models.product import Product
from app.models.product_variant import ProductVariant
from app.models.product_topping import ProductTopping
from app.api.v1.products.schemas import (
    PRODUCT_TEMPLATE_HEADERS,
    VARIANT_TEMPLATE_HEADERS,
    RECIPE_TEMPLATE_HEADERS,
    TOPPING_TEMPLATE_HEADERS,
)

_PREPARATION_LABELS = {"prepared": "Preparado", "packaged": "Empacado"}


def _bool_label(value: bool) -> str:
    """Mismo criterio que ya usa el export/import de inventario para sus columnas
    booleanas (data-model.md)."""
    return "Sí" if value else "No"


def build_products_excel(db: Session) -> io.BytesIO:
    """Arma el workbook completo en memoria, en el orden Productos → Presentaciones
    y precios → Receta → Adiciones (FR-007). Incluye todos los productos del tenant,
    activos e inactivos (Historia 1). Un producto sin receta ni adiciones solo
    genera filas en las dos primeras hojas; una carta vacía entrega las 4 hojas con
    encabezados sin filas de datos."""
    products = db.execute(
        select(Product)
        .options(
            selectinload(Product.category),
            selectinload(Product.variants).selectinload(ProductVariant.presentation),
            selectinload(Product.variants).selectinload(ProductVariant.recipe_items),
            selectinload(Product.toppings).selectinload(ProductTopping.topping_product),
        )
        .order_by(Product.name)
    ).unique().scalars().all()

    wb = Workbook()
    ws_products = wb.active
    ws_products.title = "Productos"
    ws_products.append(PRODUCT_TEMPLATE_HEADERS)

    ws_variants = wb.create_sheet("Presentaciones y precios")
    ws_variants.append(VARIANT_TEMPLATE_HEADERS)

    ws_recipe = wb.create_sheet("Receta")
    ws_recipe.append(RECIPE_TEMPLATE_HEADERS)

    ws_toppings = wb.create_sheet("Adiciones")
    ws_toppings.append(TOPPING_TEMPLATE_HEADERS)

    for product in products:
        ws_products.append([
            product.code,
            product.name,
            product.category.name if product.category else "",
            _PREPARATION_LABELS.get(product.preparation_type, product.preparation_type),
            product.description or "",
            _bool_label(product.active),
            _bool_label(product.available),
            _bool_label(product.tracks_inventory),
            _bool_label(product.is_topping),
            _bool_label(product.visible_en_qr),
        ])

        for variant in product.variants:
            ws_variants.append([
                product.code,
                variant.presentation_name,
                variant.sku or "",
                variant.name or "",
                variant.price,
                variant.display_order,
                _bool_label(variant.active),
            ])
            for recipe_item in variant.recipe_items:
                ws_recipe.append([
                    product.code,
                    variant.presentation_name,
                    recipe_item.inventory_item_name,
                    recipe_item.quantity,
                ])

        for topping in product.toppings:
            ws_toppings.append([product.code, topping.topping_product.code])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_import_errors_excel(errors: list[dict]) -> io.BytesIO:
    """Reporte descargable de las filas fallidas de un `ProductImportJob` (FR-013):
    una fila por error, con la hoja de origen, el número de fila original, su
    código (si lo trajo) y el motivo del fallo."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Hoja", "Fila", "Código", "Motivo"])
    for error in errors:
        ws.append([
            error.get("hoja") or "",
            error.get("fila"),
            error.get("codigo") or "",
            error.get("motivo") or "",
        ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
