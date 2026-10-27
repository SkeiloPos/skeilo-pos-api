from datetime import date
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.db import get_db, get_tenant
from app.core.crud import get_or_404, ensure_unique
from app.core.dependencies import get_current_user, require_tenant_admin
from app.core.models import Tenant, User
from app.core.pagination import Page, paginate
from app.core.plan_limits import require_module_access
from app.core.timezone import utc_now_naive
from app.models.inventory_item import InventoryItem
from app.models.inventory_movement import InventoryMovement
from app.models.inventory_import_job import InventoryImportJob
from app.models.unit_measure import UnitMeasure
from app.models.supplier import Supplier
from app.models.purchase import Purchase
from app.api.v1.inventory import service
from app.api.v1.inventory import supplier_items
from app.api.v1.inventory import import_ as import_service
from app.api.v1.inventory.export import build_inventory_excel, build_import_errors_excel
from app.api.v1.inventory.stock import apply_adjustment
from app.celery_task import import_inventory_task
from app.api.v1.inventory.schemas import (
    InventoryItemCreate, InventoryItemUpdate, InventoryItemResponse, InventoryItemType,
    AdjustmentIn, MovementResponse,
    SupplierCreate, SupplierUpdate, SupplierResponse,
    SupplierItemCreate, SupplierItemUpdate, SupplierItemResponse,
    PurchaseCreate, PurchaseResponse, PurchaseReceiveIn,
    LowStockResponse, InventoryImportJobResponse,
)

router = APIRouter(
    prefix="/inventory", tags=["inventory"],
    dependencies=[Depends(require_module_access("inventario"))],
)


# ============================ Insumos ============================
@router.get("/items", response_model=Page[InventoryItemResponse], summary="Listar insumos")
def list_items(
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    search: str | None = Query(None, description="Búsqueda por nombre o código (contiene, sin distinguir mayúsculas)"),
    type: InventoryItemType | None = Query(None, description="Filtrar por tipo de insumo"),
    active: bool | None = Query(None),
    low_stock: bool | None = Query(None, description="Solo insumos en o bajo el mínimo"),
    supplier_id: UUID | None = Query(None, description="Filtrar por proveedor (solo relaciones activas)"),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    stmt = service.list_items_query(
        search=search.strip() if search else None,
        type_=type.value if type is not None else None,
        active=active,
        low_stock=low_stock,
        supplier_id=supplier_id,
    )
    result = paginate(db, stmt, page, size)
    supplier_items.attach_preferred_supplier_names(db, result["items"])
    return result


@router.get("/items/low-stock", response_model=list[LowStockResponse], summary="Insumos en o bajo el mínimo")
def low_stock(
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    return db.execute(
        select(InventoryItem).where(
            InventoryItem.active.is_(True),
            InventoryItem.current_stock <= InventoryItem.min_stock,
        ).order_by(InventoryItem.name)
    ).scalars().all()


@router.get("/items/export", summary="Exportar el inventario completo a .xlsx")
def export_items(
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    items = service.list_all_items_for_export(db)
    unit_measures_by_id = {u.id: u for u in db.execute(select(UnitMeasure)).scalars().all()}
    preferred_by_item = supplier_items.preferred_relations_by_item(db, [i.id for i in items])
    buffer = build_inventory_excel(items, unit_measures_by_id, preferred_by_item)
    filename = f"inventario_{date.today().isoformat()}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _job_response(job: InventoryImportJob) -> InventoryImportJobResponse:
    # `has_errors` no es una columna mapeada (mismo patrón que
    # `preferred_supplier_name`) -- se fija aquí para que
    # `InventoryImportJobResponse` la lea vía `from_attributes`.
    job.has_errors = job.failed_count > 0
    return InventoryImportJobResponse.model_validate(job)


@router.post(
    "/items/import", response_model=InventoryImportJobResponse,
    status_code=status.HTTP_202_ACCEPTED, summary="Importar inventario desde .xlsx",
)
def import_items(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_tenant),
    user: User = Depends(require_tenant_admin),
):
    try:
        wb = load_workbook(file.file, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY, "El archivo no es un .xlsx válido"
        ) from exc

    filas_wb = list(wb.active.iter_rows(values_only=True))
    if not filas_wb:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "El archivo está vacío")
    headers, filas = filas_wb[0], [list(f) for f in filas_wb[1:]]

    problemas = import_service.validar_estructura(headers)
    if problemas:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "; ".join(problemas))

    en_curso = service.find_blocking_import_job(db)
    if en_curso is not None:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"Ya hay una importación en curso (job {en_curso.id})",
        )

    job = InventoryImportJob(
        status="procesando", file_name=file.filename or "inventario.xlsx",
        started_by_user_id=user.id, total_rows=len(filas),
    )
    db.add(job)
    db.commit()
    db.refresh(job)

    import_inventory_task.delay(tenant.schema, str(job.id), filas)
    return _job_response(job)


@router.get(
    "/items/import/current", response_model=InventoryImportJobResponse | None,
    summary="Job de importación más reciente del negocio",
)
def current_import_job(
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    """Historia 3, escenario 2: el administrador puede cerrar y volver a abrir
    la pantalla de Inventario y encontrar el resultado del import, en curso o
    ya finalizado, sin recordar ningún id."""
    job = db.execute(
        select(InventoryImportJob).order_by(InventoryImportJob.started_at.desc())
    ).scalars().first()
    return _job_response(job) if job is not None else None


def _get_import_job_or_404(db: Session, job_id: UUID) -> InventoryImportJob:
    job = db.get(InventoryImportJob, job_id)
    if job is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Import job not found")
    return job


@router.get(
    "/items/import/{job_id}", response_model=InventoryImportJobResponse,
    summary="Job de importación puntual",
)
def get_import_job(
    job_id: UUID,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    return _job_response(_get_import_job_or_404(db, job_id))


@router.post(
    "/items/import/{job_id}/cancel", response_model=InventoryImportJobResponse,
    summary="Cancelar una importación en curso",
)
def cancel_import_job(
    job_id: UUID,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    """Permite al usuario cerrar desde el front un job que quedó colgado (p.
    ej. el worker murió sin dejar el job en un estado final), sin necesitar
    acceso directo a la base de datos. Marcarlo `interrumpido` también lo
    saca de `find_blocking_import_job`, liberando el bloqueo para un import
    nuevo. No detiene el proceso del worker si sigue vivo (`procesar_archivo_
    importacion` relee `job.status` entre fila y fila y corta solo ahí)."""
    job = _get_import_job_or_404(db, job_id)
    if job.status != "procesando":
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"El job {job_id} ya no está en curso (status: {job.status})",
        )
    job.status = "interrumpido"
    job.finished_at = utc_now_naive()
    db.add(job)
    db.commit()
    db.refresh(job)
    return _job_response(job)


@router.get("/items/import/{job_id}/errors", summary="Reporte de filas fallidas de un import")
def download_import_errors(
    job_id: UUID,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    job = _get_import_job_or_404(db, job_id)
    if job.failed_count == 0:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Este import no tiene filas fallidas")

    buffer = build_import_errors_excel(job.errors)
    filename = f"errores_{job.file_name}"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/items/{item_id}", response_model=InventoryItemResponse, summary="Obtener un insumo")
def get_item(item_id: UUID, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    item = get_or_404(db, InventoryItem, item_id, "Inventory item not found")
    supplier_items.attach_preferred_supplier_names(db, [item])
    return item


@router.post("/items", response_model=InventoryItemResponse, status_code=status.HTTP_201_CREATED, summary="Crear un insumo")
def create_item(
    body: InventoryItemCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    get_or_404(db, UnitMeasure, body.unit_measure_id, "Unit measure not found")
    ensure_unique(db, InventoryItem, InventoryItem.name, body.name, "Inventory item name already exists")
    code = body.code or service.generate_next_code(db)
    ensure_unique(db, InventoryItem, InventoryItem.code, code, "Inventory item code already exists")
    item = InventoryItem(
        code=code,
        name=body.name,
        unit_measure_id=body.unit_measure_id,
        type=body.type.value,
        current_stock=body.current_stock,
        min_stock=body.min_stock,
        unit_cost=body.unit_cost,
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@router.patch("/items/{item_id}", response_model=InventoryItemResponse, summary="Actualizar un insumo")
def update_item(
    item_id: UUID,
    body: InventoryItemUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    item = get_or_404(db, InventoryItem, item_id, "Inventory item not found")
    if body.code is not None and body.code != item.code:
        ensure_unique(db, InventoryItem, InventoryItem.code, body.code, "Inventory item code already exists")
        item.code = body.code
    if body.name is not None and body.name != item.name:
        ensure_unique(db, InventoryItem, InventoryItem.name, body.name, "Inventory item name already exists")
        item.name = body.name
    if body.unit_measure_id is not None and body.unit_measure_id != item.unit_measure_id:
        get_or_404(db, UnitMeasure, body.unit_measure_id, "Unit measure not found")
        has_movements = db.execute(
            select(InventoryMovement.id)
            .where(InventoryMovement.inventory_item_id == item.id)
            .limit(1)
        ).first() is not None
        if has_movements:
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                "No se puede cambiar la unidad base de un insumo con movimientos registrados",
            )
        item.unit_measure_id = body.unit_measure_id
    if body.type is not None:
        item.type = body.type.value
    if body.min_stock is not None:
        item.min_stock = body.min_stock
    if body.unit_cost is not None:
        item.unit_cost = body.unit_cost
    if body.active is not None:
        item.active = body.active
    db.commit()
    db.refresh(item)
    return item


@router.post("/items/{item_id}/adjust", response_model=MovementResponse, summary="Ajustar stock (delta con signo)")
def adjust_item(
    item_id: UUID,
    body: AdjustmentIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_tenant_admin),
):
    get_or_404(db, InventoryItem, item_id, "Inventory item not found")
    movement = apply_adjustment(db, item_id, signed_delta=body.signed_delta, reason=body.reason, user_id=user.id)
    db.commit()
    db.refresh(movement)
    return movement


@router.get("/items/{item_id}/movements", response_model=Page[MovementResponse], summary="Kardex de un insumo")
def item_movements(
    item_id: UUID,
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    get_or_404(db, InventoryItem, item_id, "Inventory item not found")
    stmt = (
        select(InventoryMovement)
        .where(InventoryMovement.inventory_item_id == item_id)
        .order_by(InventoryMovement.moved_at.desc())
    )
    return paginate(db, stmt, page, size)


# ============================ Proveedores ============================
@router.get("/suppliers", response_model=list[SupplierResponse], summary="Listar proveedores")
def list_suppliers(
    exclude_system: bool = Query(False, description="Excluye proveedores de sistema (FR-012)"),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    stmt = select(Supplier)
    if exclude_system:
        stmt = stmt.where(Supplier.is_system.is_(False))
    return db.execute(stmt.order_by(Supplier.name)).scalars().all()


@router.post("/suppliers", response_model=SupplierResponse, status_code=status.HTTP_201_CREATED, summary="Crear proveedor")
def create_supplier(body: SupplierCreate, db: Session = Depends(get_db), _: User = Depends(require_tenant_admin)):
    supplier = Supplier(**body.model_dump())
    db.add(supplier)
    db.commit()
    db.refresh(supplier)
    return supplier


@router.patch("/suppliers/{supplier_id}", response_model=SupplierResponse, summary="Actualizar proveedor")
def update_supplier(
    supplier_id: UUID,
    body: SupplierUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    supplier = get_or_404(db, Supplier, supplier_id, "Supplier not found")
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(supplier, k, v)
    db.commit()
    db.refresh(supplier)
    return supplier


# ===================== Relación proveedor-insumo (spec 106) =====================
@router.get(
    "/items/{item_id}/suppliers", response_model=list[SupplierItemResponse],
    summary="Proveedores que venden este insumo",
)
def list_item_suppliers(
    item_id: UUID,
    include_inactive: bool = Query(False),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    relations = supplier_items.list_for_item(db, item_id, include_inactive)
    return [supplier_items.to_response(db, r) for r in relations]


@router.post(
    "/items/{item_id}/suppliers", response_model=SupplierItemResponse,
    summary="Vincular un proveedor a este insumo",
)
def create_item_supplier(
    item_id: UUID,
    body: SupplierItemCreate,
    response: Response,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    relation, status_code = supplier_items.create_or_reactivate(
        db, supplier_id=body.supplier_id, inventory_item_id=item_id, body=body,
    )
    response.status_code = status_code
    return supplier_items.to_response(db, relation)


@router.get(
    "/suppliers/{supplier_id}/items", response_model=list[SupplierItemResponse],
    summary="Insumos que vende este proveedor",
)
def list_supplier_items(
    supplier_id: UUID,
    include_inactive: bool = Query(False),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    relations = supplier_items.list_for_supplier(db, supplier_id, include_inactive)
    return [supplier_items.to_response(db, r) for r in relations]


@router.post(
    "/suppliers/{supplier_id}/items", response_model=SupplierItemResponse,
    summary="Vincular este proveedor a un insumo",
)
def create_supplier_item(
    supplier_id: UUID,
    body: SupplierItemCreate,
    response: Response,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    relation, status_code = supplier_items.create_or_reactivate(
        db, supplier_id=supplier_id, inventory_item_id=body.inventory_item_id, body=body,
    )
    response.status_code = status_code
    return supplier_items.to_response(db, relation)


@router.patch(
    "/supplier-items/{relation_id}", response_model=SupplierItemResponse,
    summary="Editar, preferir o (des)activar una relación proveedor-insumo",
)
def update_supplier_item(
    relation_id: UUID,
    body: SupplierItemUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_tenant_admin),
):
    relation, was_preferred = supplier_items.update(db, relation_id, body)
    return supplier_items.to_response(db, relation, was_preferred=was_preferred)


# ============================ Compras ============================
@router.get(
    "/purchases", response_model=Page[PurchaseResponse], summary="Listar compras",
    dependencies=[Depends(require_module_access("compras"))],
)
def list_purchases(
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    stmt = select(Purchase).order_by(Purchase.purchased_at.desc())
    result = paginate(db, stmt, page, size)
    for p in result["items"]:
        supplier_items.attach_conversion_factors(db, p.items)
    return result


@router.post(
    "/purchases", response_model=PurchaseResponse, status_code=status.HTTP_201_CREATED,
    summary="Registrar compra (da alta de stock total)",
    dependencies=[Depends(require_module_access("compras"))],
)
def create_purchase(
    body: PurchaseCreate,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_tenant),
    user: User = Depends(require_tenant_admin),
):
    purchase = service.create_purchase(db, body, user_id=user.id, tenant=tenant)
    supplier_items.attach_conversion_factors(db, purchase.items)
    return purchase


@router.post(
    "/purchases/order", response_model=PurchaseResponse, status_code=status.HTTP_201_CREATED,
    summary="Crear orden de compra (draft, sin alta de stock) — RF-022",
    dependencies=[Depends(require_module_access("compras"))],
)
def create_purchase_order(
    body: PurchaseCreate,
    db: Session = Depends(get_db),
    tenant: Tenant = Depends(get_tenant),
    user: User = Depends(require_tenant_admin),
):
    purchase = service.create_purchase_order(db, body, user_id=user.id, tenant=tenant)
    supplier_items.attach_conversion_factors(db, purchase.items)
    return purchase


@router.post(
    "/purchases/{purchase_id}/receive", response_model=PurchaseResponse,
    summary="Recibir una orden de compra (parcial o total) — RF-022",
    dependencies=[Depends(require_module_access("compras"))],
)
def receive_purchase(
    purchase_id: UUID,
    body: PurchaseReceiveIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_tenant_admin),
):
    purchase = service.receive_purchase(db, purchase_id, body, user_id=user.id)
    supplier_items.attach_conversion_factors(db, purchase.items)
    return purchase


@router.post(
    "/purchases/{purchase_id}/void", response_model=PurchaseResponse,
    summary="Anular una compra confirmada (spec 134)",
    dependencies=[Depends(require_module_access("compras"))],
)
def void_purchase(
    purchase_id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(require_tenant_admin),
):
    purchase = service.void_purchase(db, purchase_id, user_id=user.id)
    supplier_items.attach_conversion_factors(db, purchase.items)
    return purchase
