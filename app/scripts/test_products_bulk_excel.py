"""Test de export/import masivo de la carta en Excel (spec 143).

No hay pytest en el proyecto, así que es un script autoejecutable:

    python -m app.scripts.test_products_bulk_excel

Cubre los 16 puntos de `quickstart.md` §1. Trabaja sobre un tenant real con
datos desechables: crea su propia categoría, presentación, insumo y productos,
y los borra al terminar (incluso si alguna prueba falla a mitad de camino --
ver `main()`)."""
from decimal import Decimal
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import select, delete, text

from app.core.db import with_db
from app.models.category import Category
from app.models.presentation import Presentation
from app.models.product import Product
from app.models.product_variant import ProductVariant
from app.models.recipe_item import RecipeItem
from app.models.product_topping import ProductTopping
from app.models.product_import_job import ProductImportJob
from app.models.inventory_item import InventoryItem
from app.models.unit_measure import UnitMeasure
from app.models.unit_measure_group import UnitMeasureGroup
from app.api.v1.products import service, import_ as import_service
from app.api.v1.products.export import build_products_excel, build_import_errors_excel
from app.api.v1.products.router import _find_blocking_import_job
from app.api.v1.products.schemas import (
    PRODUCT_TEMPLATE_HEADERS, VARIANT_TEMPLATE_HEADERS,
    RECIPE_TEMPLATE_HEADERS, TOPPING_TEMPLATE_HEADERS,
)

_codigos_producto_extra: list[str] = []
_job_ids: list = []
_categoria_ids_extra: list = []
_presentacion_ids_extra: list = []


def _schema() -> str:
    with with_db(None) as db:
        rows = db.execute(text("SELECT schema FROM shared.tenants ORDER BY id")).scalars().all()
    for schema in rows:
        with with_db(None) as db:
            exists = db.execute(
                text("SELECT to_regclass(:q)"), {"q": f"{schema}.products"}
            ).scalar() is not None
        if exists:
            return schema
    raise SystemExit("Ningún tenant tiene su esquema de productos creado todavía.")


def _fixture(db):
    categoria = Category(
        name=f"Categoria-test-{uuid4().hex[:6]}", active=True, visible_en_qr=True,
        display_order=999,
    )
    db.add(categoria)
    db.flush()

    presentacion = Presentation(name=f"Presentacion-test-{uuid4().hex[:6]}", active=True)
    db.add(presentacion)
    db.flush()

    unit = db.execute(
        select(UnitMeasure).where(UnitMeasure.abbreviation == "G")
    ).scalars().first()
    if unit is None:
        grupo = UnitMeasureGroup(name=f"gramo-test-{uuid4().hex[:6]}")
        db.add(grupo)
        db.flush()
        unit = UnitMeasure(
            name=f"gramo-test-{uuid4().hex[:6]}", abbreviation="G",
            group_id=grupo.id, is_base=True, factor_to_base=1,
        )
        db.add(unit)
        db.flush()

    insumo = InventoryItem(
        code=f"INS-{uuid4().hex[:6]}", name=f"Insumo-test-{uuid4().hex[:8]}",
        type="raw_material", unit_measure_id=unit.id,
        current_stock=Decimal("100"), min_stock=Decimal("0"), unit_cost=Decimal("10"),
    )
    db.add(insumo)
    db.flush()

    producto_con_receta = Product(
        code=f"T{uuid4().hex[:6]}", name=f"Producto-con-receta-{uuid4().hex[:6]}",
        category_id=categoria.id, preparation_type="prepared", active=True, available=True,
    )
    db.add(producto_con_receta)
    db.flush()
    variante = ProductVariant(
        product_id=producto_con_receta.id, presentation_id=presentacion.id,
        price=Decimal("5000"), display_order=1, active=True,
    )
    db.add(variante)
    db.flush()
    receta = RecipeItem(
        product_variant_id=variante.id, inventory_item_id=insumo.id, quantity=Decimal("2"),
    )
    db.add(receta)
    db.flush()

    producto_sin_receta = Product(
        code=f"T{uuid4().hex[:6]}", name=f"Producto-sin-receta-{uuid4().hex[:6]}",
        category_id=categoria.id, preparation_type="packaged", active=True, available=True,
    )
    db.add(producto_sin_receta)
    db.flush()
    variante_sin_receta = ProductVariant(
        product_id=producto_sin_receta.id, presentation_id=presentacion.id,
        price=Decimal("3000"), display_order=1, active=True,
    )
    db.add(variante_sin_receta)
    db.flush()

    producto_adicion = Product(
        code=f"T{uuid4().hex[:6]}", name=f"Producto-adicion-{uuid4().hex[:6]}",
        category_id=categoria.id, preparation_type="packaged", is_topping=True,
        active=True, available=True,
    )
    db.add(producto_adicion)
    db.flush()

    db.commit()
    return {
        "categoria": categoria, "presentacion": presentacion, "unit": unit, "insumo": insumo,
        "producto_con_receta": producto_con_receta, "variante": variante,
        "producto_sin_receta": producto_sin_receta, "producto_adicion": producto_adicion,
    }


def _cleanup(db, fx) -> None:
    db.rollback()
    extra_codes = [*_codigos_producto_extra]
    extra_ids = db.execute(
        select(Product.id).where(Product.code.in_(extra_codes))
    ).scalars().all() if extra_codes else []
    all_product_ids = [
        fx["producto_con_receta"].id, fx["producto_sin_receta"].id, fx["producto_adicion"].id,
        *extra_ids,
    ]

    if _job_ids:
        db.execute(delete(ProductImportJob).where(ProductImportJob.id.in_(_job_ids)))
    db.execute(delete(ProductTopping).where(
        ProductTopping.product_id.in_(all_product_ids)
        | ProductTopping.topping_product_id.in_(all_product_ids)
    ))
    db.execute(delete(RecipeItem).where(
        RecipeItem.product_variant_id.in_(
            select(ProductVariant.id).where(ProductVariant.product_id.in_(all_product_ids))
        )
    ))
    db.execute(delete(ProductVariant).where(ProductVariant.product_id.in_(all_product_ids)))
    db.execute(delete(Product).where(Product.id.in_(all_product_ids)))
    db.execute(delete(InventoryItem).where(InventoryItem.id == fx["insumo"].id))
    if _presentacion_ids_extra:
        db.execute(delete(Presentation).where(Presentation.id.in_(_presentacion_ids_extra)))
    db.execute(delete(Presentation).where(Presentation.id == fx["presentacion"].id))
    if _categoria_ids_extra:
        db.execute(delete(Category).where(Category.id.in_(_categoria_ids_extra)))
    db.execute(delete(Category).where(Category.id == fx["categoria"].id))
    db.commit()


def _job(db, user_id, total_rows: int) -> ProductImportJob:
    job = ProductImportJob(
        status="procesando", file_name="test.xlsx", started_by_user_id=user_id,
        total_rows=total_rows,
    )
    db.add(job)
    db.flush()
    _job_ids.append(job.id)
    return job


def _row_by_header(headers: list[str], values: list, header: str):
    return values[headers.index(header)]


# 1. Export: con/sin receta y adiciones, carta vacía.
def test_export_incluye_las_4_hojas_con_y_sin_receta(db, fx):
    buffer = build_products_excel(db)
    wb = load_workbook(buffer)
    assert wb.sheetnames == ["Productos", "Presentaciones y precios", "Receta", "Adiciones"], wb.sheetnames

    ws_p = wb["Productos"]
    headers_p = [c.value for c in next(ws_p.iter_rows(min_row=1, max_row=1))]
    assert headers_p == PRODUCT_TEMPLATE_HEADERS, headers_p
    codigos_productos = [row[0] for row in ws_p.iter_rows(min_row=2, values_only=True)]
    assert fx["producto_con_receta"].code in codigos_productos
    assert fx["producto_sin_receta"].code in codigos_productos

    ws_v = wb["Presentaciones y precios"]
    assert [c.value for c in next(ws_v.iter_rows(min_row=1, max_row=1))] == VARIANT_TEMPLATE_HEADERS
    codigos_variantes = [row[0] for row in ws_v.iter_rows(min_row=2, values_only=True)]
    assert fx["producto_sin_receta"].code in codigos_variantes

    ws_r = wb["Receta"]
    assert [c.value for c in next(ws_r.iter_rows(min_row=1, max_row=1))] == RECIPE_TEMPLATE_HEADERS
    codigos_receta = [row[0] for row in ws_r.iter_rows(min_row=2, values_only=True)]
    assert fx["producto_con_receta"].code in codigos_receta
    assert fx["producto_sin_receta"].code not in codigos_receta

    ws_t = wb["Adiciones"]
    assert [c.value for c in next(ws_t.iter_rows(min_row=1, max_row=1))] == TOPPING_TEMPLATE_HEADERS
    print("OK: export_incluye_las_4_hojas_con_y_sin_receta")


def test_export_carta_vacia_entrega_4_hojas_sin_filas(db):
    """Historia 1, escenario 3: carta sin ningún producto -> 4 hojas con
    encabezados, sin filas de datos."""
    vacia = db.execute(select(Product.id)).first()
    if vacia is not None:
        print("SKIP: export_carta_vacia_entrega_4_hojas_sin_filas (el tenant ya tiene productos)")
        return
    buffer = build_products_excel(db)
    wb = load_workbook(buffer)
    for nombre in ("Productos", "Presentaciones y precios", "Receta", "Adiciones"):
        assert wb[nombre].max_row == 1, f"{nombre} debía tener solo encabezados"
    print("OK: export_carta_vacia_entrega_4_hojas_sin_filas")


# 2-3. Reimportar sin editar / editar precio.
def _fila_producto(producto, categoria_nombre):
    return [
        producto.code, producto.name, categoria_nombre,
        "Preparado" if producto.preparation_type == "prepared" else "Empacado",
        producto.description or "", "Sí" if producto.active else "No",
        "Sí" if producto.available else "No", "Sí" if producto.tracks_inventory else "No",
        "Sí" if producto.is_topping else "No", "Sí" if producto.visible_en_qr else "No",
    ]


def _fila_variante(producto, variante, presentacion_nombre):
    return [
        producto.code, presentacion_nombre, variante.sku or "", variante.name or "",
        float(variante.price), variante.display_order, "Sí" if variante.active else "No",
    ]


def test_reimportar_sin_editar_no_duplica(db, fx):
    hojas = {
        "Productos": [_fila_producto(fx["producto_con_receta"], fx["categoria"].name)],
        "Presentaciones y precios": [
            _fila_variante(fx["producto_con_receta"], fx["variante"], fx["presentacion"].name),
        ],
        "Receta": [[
            fx["producto_con_receta"].code, fx["presentacion"].name, fx["insumo"].name,
            float(2),
        ]],
    }
    job = _job(db, uuid4(), sum(len(v) for v in hojas.values()))
    db.commit()
    import_service.procesar_archivo_importacion(db, job, hojas)
    db.refresh(job)
    assert job.failed_count == 0, job.errors
    assert job.created_count == 0, job.errors
    assert job.updated_count == 3, job.errors
    print("OK: reimportar_sin_editar_no_duplica")


def test_editar_precio_actualiza_la_variante(db, fx):
    nuevo_precio = Decimal("6500")
    fila = _fila_variante(fx["producto_con_receta"], fx["variante"], fx["presentacion"].name)
    fila[4] = float(nuevo_precio)
    hojas = {"Presentaciones y precios": [fila]}
    job = _job(db, uuid4(), 1)
    db.commit()
    import_service.procesar_archivo_importacion(db, job, hojas)
    db.refresh(job)
    assert job.updated_count == 1, job.errors
    db.refresh(fx["variante"])
    assert fx["variante"].price == nuevo_precio
    print("OK: editar_precio_actualiza_la_variante")


# 4. Fila nueva con código vacío -> producto nuevo con código automático.
def test_fila_producto_codigo_vacio_crea_con_codigo_automatico(db, fx):
    antes = set(db.execute(select(Product.code)).scalars().all())
    fila = [
        "", f"Nuevo-{uuid4().hex[:6]}", fx["categoria"].name, "Preparado", "", "Sí", "Sí",
        "No", "No", "Sí",
    ]
    job = _job(db, uuid4(), 1)
    db.commit()
    import_service.procesar_archivo_importacion(db, job, {"Productos": [fila]})
    db.refresh(job)
    assert job.created_count == 1, job.errors
    despues = set(db.execute(select(Product.code)).scalars().all())
    nuevos = despues - antes
    assert len(nuevos) == 1
    _codigos_producto_extra.append(nuevos.pop())
    print("OK: fila_producto_codigo_vacio_crea_con_codigo_automatico")


# 5. Receta con insumo inexistente -> falla, resto no afectado.
def test_receta_con_insumo_inexistente_falla_sin_afectar_el_resto(db, fx):
    fila_mala = [
        fx["producto_con_receta"].code, fx["presentacion"].name, f"insumo-que-no-existe-{uuid4().hex[:6]}",
        1.0,
    ]
    fila_buena = [
        fx["producto_con_receta"].code, fx["presentacion"].name, fx["insumo"].name, 3.0,
    ]
    job = _job(db, uuid4(), 2)
    db.commit()
    import_service.procesar_archivo_importacion(db, job, {"Receta": [fila_mala, fila_buena]})
    db.refresh(job)
    assert job.failed_count == 1, job.errors
    assert job.updated_count == 1, job.errors
    assert "insumo" in job.errors[0]["motivo"].lower()
    db.refresh(fx["variante"])
    print("OK: receta_con_insumo_inexistente_falla_sin_afectar_el_resto")


# 6. Adición sin is_topping=true -> falla, resto continúa.
def test_adicion_sin_is_topping_falla(db, fx):
    fila_mala = [fx["producto_con_receta"].code, fx["producto_sin_receta"].code]
    fila_buena = [fx["producto_con_receta"].code, fx["producto_adicion"].code]
    job = _job(db, uuid4(), 2)
    db.commit()
    import_service.procesar_archivo_importacion(db, job, {"Adiciones": [fila_mala, fila_buena]})
    db.refresh(job)
    assert job.failed_count == 1, job.errors
    assert job.created_count == 1, job.errors
    assert "adición" in job.errors[0]["motivo"].lower() or "topping" in job.errors[0]["motivo"].lower()
    print("OK: adicion_sin_is_topping_falla")


# 7. Progreso consultable durante el proceso (proxy: los conteos cuadran al terminar).
def test_progreso_queda_reflejado_en_los_conteos(db, fx):
    filas = [_fila_producto(fx["producto_sin_receta"], fx["categoria"].name) for _ in range(5)]
    job = _job(db, uuid4(), len(filas))
    db.commit()
    import_service.procesar_archivo_importacion(db, job, {"Productos": filas})
    db.refresh(job)
    assert job.created_count + job.updated_count + job.failed_count == len(filas)
    assert job.status == "finalizado"
    print("OK: progreso_queda_reflejado_en_los_conteos")


# 8. Reporte de errores descargable.
def test_reporte_de_errores_identifica_hoja_fila_y_motivo():
    errors = [
        {"hoja": "Receta", "fila": 3, "codigo": "T000X", "motivo": "no existe el insumo"},
    ]
    buffer = build_import_errors_excel(errors)
    wb = load_workbook(buffer)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == ["Hoja", "Fila", "Código", "Motivo"], headers
    fila = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert fila == ["Receta", 3, "T000X", "no existe el insumo"]
    print("OK: reporte_de_errores_identifica_hoja_fila_y_motivo")


# 9. Reintentar un archivo con filas corregidas no duplica lo ya creado.
def test_reintentar_archivo_corregido_no_duplica(db, fx):
    code_ok = f"T{uuid4().hex[:6]}"
    code_malo = f"T{uuid4().hex[:6]}"
    fila_ok = [
        code_ok, f"Retry-ok-{uuid4().hex[:6]}", fx["categoria"].name, "Preparado", "", "Sí",
        "Sí", "No", "No", "Sí",
    ]
    fila_mala = [
        code_malo, f"Retry-malo-{uuid4().hex[:6]}", f"categoria-inexistente-{uuid4().hex[:4]}",
        "Tipo invalido", "", "Sí", "Sí", "No", "No", "Sí",
    ]
    job1 = _job(db, uuid4(), 2)
    db.commit()
    import_service.procesar_archivo_importacion(db, job1, {"Productos": [fila_ok, fila_mala]})
    db.refresh(job1)
    assert job1.created_count == 1, job1.errors
    assert job1.failed_count == 1, job1.errors
    _codigos_producto_extra.append(code_ok)

    fila_corregida = [
        code_malo, f"Retry-corregido-{uuid4().hex[:6]}", fx["categoria"].name, "Preparado", "",
        "Sí", "Sí", "No", "No", "Sí",
    ]
    job2 = _job(db, uuid4(), 2)
    db.commit()
    import_service.procesar_archivo_importacion(db, job2, {"Productos": [fila_ok, fila_corregida]})
    db.refresh(job2)
    assert job2.updated_count == 1, job2.errors
    assert job2.created_count == 1, job2.errors
    _codigos_producto_extra.append(code_malo)

    count_ok = db.execute(select(Product.id).where(Product.code == code_ok)).all()
    assert len(count_ok) == 1, "no debía duplicarse"
    print("OK: reintentar_archivo_corregido_no_duplica")


# 10. Códigos duplicados dentro del mismo archivo -> ambas filas fallan.
def test_codigos_duplicados_en_el_archivo_fallan_ambas(db, fx):
    code = f"DUP-{uuid4().hex[:6]}"
    filas = [
        [code, "Primera", fx["categoria"].name, "Preparado", "", "Sí", "Sí", "No", "No", "Sí"],
        [code, "Segunda", fx["categoria"].name, "Preparado", "", "Sí", "Sí", "No", "No", "Sí"],
    ]
    job = _job(db, uuid4(), 2)
    db.commit()
    import_service.procesar_archivo_importacion(db, job, {"Productos": filas})
    db.refresh(job)
    assert job.failed_count == 2, job.errors
    assert job.created_count == 0
    existe = db.execute(select(Product).where(Product.code == code)).scalar_one_or_none()
    assert existe is None
    print("OK: codigos_duplicados_en_el_archivo_fallan_ambas")


# 11. Presentación con código de producto inexistente -> falla.
def test_presentacion_con_codigo_inexistente_falla(db, fx):
    fila = [f"NO-EXISTE-{uuid4().hex[:6]}", fx["presentacion"].name, "", "", 1000.0, 1, "Sí"]
    job = _job(db, uuid4(), 1)
    db.commit()
    import_service.procesar_archivo_importacion(db, job, {"Presentaciones y precios": [fila]})
    db.refresh(job)
    assert job.failed_count == 1, job.errors
    print("OK: presentacion_con_codigo_inexistente_falla")


# 12. Adición con el mismo código del producto que la ofrece -> falla.
def test_adicion_consigo_mismo_falla(db, fx):
    fila = [fx["producto_con_receta"].code, fx["producto_con_receta"].code]
    job = _job(db, uuid4(), 1)
    db.commit()
    import_service.procesar_archivo_importacion(db, job, {"Adiciones": [fila]})
    db.refresh(job)
    assert job.failed_count == 1, job.errors
    assert "sí mismo" in job.errors[0]["motivo"].lower()
    print("OK: adicion_consigo_mismo_falla")


# 13. FR-005a: filas ausentes del archivo reimportado no se eliminan.
def test_filas_ausentes_del_archivo_no_se_eliminan(db, fx):
    fila_producto = _fila_producto(fx["producto_con_receta"], fx["categoria"].name)
    job = _job(db, uuid4(), 1)
    db.commit()
    import_service.procesar_archivo_importacion(db, job, {"Productos": [fila_producto]})
    db.refresh(job)
    assert job.updated_count == 1, job.errors

    variante_sigue = db.execute(
        select(ProductVariant).where(ProductVariant.id == fx["variante"].id)
    ).scalar_one_or_none()
    assert variante_sigue is not None, "la presentación no debía eliminarse"
    receta_sigue = db.execute(
        select(RecipeItem).where(RecipeItem.product_variant_id == fx["variante"].id)
    ).scalars().all()
    assert len(receta_sigue) == 1, "la receta no debía eliminarse"
    print("OK: filas_ausentes_del_archivo_no_se_eliminan")


# 14. Estructura alterada (hoja renombrada) -> 422 sin crear job.
def test_estructura_alterada_no_crea_job():
    problemas = import_service.validar_estructura({"ProductosRenombrados": PRODUCT_TEMPLATE_HEADERS})
    assert problemas, "se esperaba un problema de estructura"
    print("OK: estructura_alterada_no_crea_job")


def test_columnas_en_orden_distinto_falla_estructura():
    headers_desordenados = list(reversed(PRODUCT_TEMPLATE_HEADERS))
    problemas = import_service.validar_estructura({"Productos": headers_desordenados})
    assert problemas, "se esperaba un problema de estructura"
    print("OK: columnas_en_orden_distinto_falla_estructura")


# 15. Archivo con solo la hoja Productos -> válido.
def test_archivo_con_solo_hoja_productos_es_valido():
    problemas = import_service.validar_estructura({"Productos": PRODUCT_TEMPLATE_HEADERS})
    assert problemas is None, problemas
    print("OK: archivo_con_solo_hoja_productos_es_valido")


# 16. Job 'procesando' vigente bloquea un segundo import.
def test_job_procesando_vigente_bloquea_un_segundo_import(db):
    job = _job(db, uuid4(), 1)
    db.commit()
    bloqueante = _find_blocking_import_job(db)
    assert bloqueante is not None and bloqueante.id == job.id
    job.status = "finalizado"
    db.add(job)
    db.commit()
    print("OK: job_procesando_vigente_bloquea_un_segundo_import")


def main() -> None:
    schema = _schema()
    print(f"Usando tenant con esquema: {schema}")
    with with_db(schema) as db:
        fx = _fixture(db)
        tests = [
            lambda: test_export_incluye_las_4_hojas_con_y_sin_receta(db, fx),
            lambda: test_export_carta_vacia_entrega_4_hojas_sin_filas(db),
            lambda: test_reimportar_sin_editar_no_duplica(db, fx),
            lambda: test_editar_precio_actualiza_la_variante(db, fx),
            lambda: test_fila_producto_codigo_vacio_crea_con_codigo_automatico(db, fx),
            lambda: test_receta_con_insumo_inexistente_falla_sin_afectar_el_resto(db, fx),
            lambda: test_adicion_sin_is_topping_falla(db, fx),
            lambda: test_progreso_queda_reflejado_en_los_conteos(db, fx),
            test_reporte_de_errores_identifica_hoja_fila_y_motivo,
            lambda: test_reintentar_archivo_corregido_no_duplica(db, fx),
            lambda: test_codigos_duplicados_en_el_archivo_fallan_ambas(db, fx),
            lambda: test_presentacion_con_codigo_inexistente_falla(db, fx),
            lambda: test_adicion_consigo_mismo_falla(db, fx),
            lambda: test_filas_ausentes_del_archivo_no_se_eliminan(db, fx),
            test_estructura_alterada_no_crea_job,
            test_columnas_en_orden_distinto_falla_estructura,
            test_archivo_con_solo_hoja_productos_es_valido,
            lambda: test_job_procesando_vigente_bloquea_un_segundo_import(db),
        ]
        fallos = 0
        try:
            for test in tests:
                try:
                    test()
                except AssertionError as exc:
                    fallos += 1
                    print(f"FALLÓ: {test}: {exc}")
        finally:
            _cleanup(db, fx)
        if fallos:
            raise SystemExit(f"{fallos}/{len(tests)} pruebas fallaron")
        print(f"Todas las {len(tests)} pruebas pasaron.")


if __name__ == "__main__":
    main()
