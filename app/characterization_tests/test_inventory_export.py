"""Tests de la funcionalidad de export/import de inventario — spec 086 (origen),
ampliada por spec 109-inventario-excel-masivo (A-123): GET /inventory/items/export
(Historia 1 de la 086: descarga del reporte completo; Historia 2: consistencia de
columnas y datos exportados; Historia 1 de la 109: plantilla única de export/import).

No son characterization tests -- el endpoint no existe antes de la spec 086
(comportamiento enteramente nuevo, ver `research.md`/`contracts/
inventory-export-api.md` de esa spec). Por eso esta spec 109 actualiza sus
aserciones (nuevas columnas) directamente, sin necesitar la autorización especial
del Principio III (ver A-123, plan.md de la 109). Combina las tablas de
`plan_fixtures.py` (tenants/plans/roles/users, para el gating de plan y el
rol ADMIN) con las de `fixtures.py` (insumos/unidades de medida), mismo
patrón que `test_plan_gating_inventory_fields.py`.

Ejecutar solo este módulo:

    python -m unittest app.characterization_tests.test_inventory_export -v
"""
from __future__ import annotations

import io
import unittest
from datetime import timedelta
from decimal import Decimal
from unittest import mock
from uuid import UUID

import openpyxl
from fastapi import FastAPI
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from starlette.testclient import TestClient

from app.characterization_tests import fixtures as f
from app.characterization_tests import plan_fixtures as pf
from app.api.v1.inventory.router import router as inventory_router
from app.api.v1.inventory import import_ as import_service
from app.core.db import get_db, get_tenant
from app.core.dependencies import get_current_user
from app.core.models import Base
from app.core.timezone import utc_now
from app.models.inventory_import_job import InventoryImportJob
from app.models.inventory_item import InventoryItem
from app.models.supplier import Supplier
from app.models.supplier_item import SupplierItem
from app.models.unit_measure import UnitMeasure
from app.models.unit_measure_group import UnitMeasureGroup
from app.api.v1.inventory.schemas import INVENTORY_TEMPLATE_HEADERS

_TABLE_NAMES = [
    # plan_fixtures.py — gating de plan y rol ADMIN
    "tenants", "plans", "roles", "users",
    # fixtures.py — insumos y unidades de medida
    "inventory_items", "inventory_movements", "unit_measures", "unit_measure_groups",
    # spec 109: el export siempre consulta la relación preferida, aunque esté vacía;
    # el import necesita además la tabla de jobs.
    "suppliers", "supplier_items", "inventory_import_jobs",
]


def _uid():
    import uuid
    return uuid.uuid4()


def _new_session() -> Session:
    """`check_same_thread=False`: a diferencia de las demás fixtures del
    paquete, este módulo ejercita el endpoint vía `TestClient` (spec 086),
    que ejecuta el handler en un thread distinto al del test — mismo ajuste
    que `test_notifications.py` (spec 077) ya necesitó por el mismo motivo."""
    tables = [t for t in Base.metadata.tables.values() if t.name in _TABLE_NAMES]
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    conn = engine.connect().execution_options(
        schema_translate_map={"tenant": None, "shared": None}
    )
    Base.metadata.create_all(bind=conn, tables=tables)
    conn.commit()
    return Session(bind=conn)


def _build_client(db: Session, tenant, user) -> TestClient:
    """App mínima con el router real de inventario -- mismo patrón que
    `test_notifications.py` (spec 077): se sobreescriben `get_db`,
    `get_current_user` y `get_tenant` (esta última la usa
    `require_module_access("inventario")`, aplicado a nivel de router)."""
    app = FastAPI()
    app.include_router(inventory_router, prefix="/api/v1")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[get_tenant] = lambda: tenant
    return TestClient(app, raise_server_exceptions=False)


def _con_modulo_inventario(db: Session):
    plan = pf.make_plan(db, inventario_access=True)
    tenant = pf.make_tenant(db, plan=plan)
    return tenant


class ExportEndpointTests(unittest.TestCase):
    """Historia 1 (P1): descarga del reporte completo."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)

    def test_admin_con_insumos_descarga_n_mas_1_filas(self):
        for _ in range(3):
            f.make_inventory_item(self.db)
        f.make_inventory_item(self.db, active=False)
        self.db.commit()

        admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")
        self.db.commit()
        client = _build_client(self.db, self.tenant, admin)

        resp = client.get("/api/v1/inventory/items/export")

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="inventario_', resp.headers["content-disposition"])
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.active.max_row, 4 + 1)  # 4 insumos (3 activos + 1 inactivo)

    def test_inventario_vacio_descarga_solo_el_encabezado(self):
        admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")
        self.db.commit()
        client = _build_client(self.db, self.tenant, admin)

        resp = client.get("/api/v1/inventory/items/export")

        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.active.max_row, 1)

    def test_usuario_sin_rol_admin_recibe_403_sin_archivo(self):
        cashier = pf.make_user(self.db, self.tenant, role_name="CASHIER")
        self.db.commit()
        client = _build_client(self.db, self.tenant, cashier)

        resp = client.get("/api/v1/inventory/items/export")

        self.assertEqual(resp.status_code, 403)
        self.assertNotIn("content-disposition", resp.headers)

    def test_fallo_a_mitad_de_generacion_responde_500_sin_archivo_parcial(self):
        admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")
        self.db.commit()
        client = _build_client(self.db, self.tenant, admin)

        with mock.patch(
            "app.api.v1.inventory.router.build_inventory_excel",
            side_effect=RuntimeError("fallo simulado a mitad de la generación"),
        ):
            resp = client.get("/api/v1/inventory/items/export")

        self.assertEqual(resp.status_code, 500)
        self.assertNotIn("content-disposition", resp.headers)
        self.assertNotIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resp.headers.get("content-type", ""),
        )


class ExportColumnsTests(unittest.TestCase):
    """Historia 2 (P2): encabezados exactos y celdas numéricas nativas."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)
        self.admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")

    def _export(self):
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get("/api/v1/inventory/items/export")
        self.assertEqual(resp.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(resp.content))

    def test_fila_1_son_los_encabezados_exactos_en_orden(self):
        """spec 138, Historia 1: presentación, factor de conversión y último
        costo del proveedor preferido ya no forman parte de la plantilla (solo
        son editables desde el módulo de proveedores), pero el nombre del
        proveedor preferido sí se conserva."""
        wb = self._export()
        headers = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers, [
            "Código", "Nombre", "Tipo", "Unidad", "Stock", "Mínimo", "Costo",
            "Proveedor preferido",
        ])

    def test_celdas_stock_minimo_costo_son_numericas_no_texto(self):
        f.make_inventory_item(
            self.db, current_stock=Decimal("12.500"), min_stock=Decimal("3.000"),
            unit_cost=Decimal("15.00"),
        )
        wb = self._export()
        row = next(wb.active.iter_rows(min_row=2, max_row=2))
        stock_cell, min_cell, cost_cell = row[4], row[5], row[6]
        self.assertEqual(stock_cell.data_type, "n")
        self.assertEqual(min_cell.data_type, "n")
        self.assertEqual(cost_cell.data_type, "n")

    def test_costo_15_exporta_numero_15_no_texto_con_simbolo_de_moneda(self):
        f.make_inventory_item(self.db, unit_cost=Decimal("15.00"))
        wb = self._export()
        row = next(wb.active.iter_rows(min_row=2, max_row=2))
        cost_cell = row[6]
        self.assertEqual(float(cost_cell.value), 15.0)
        self.assertNotIsInstance(cost_cell.value, str)

    def test_codigo_es_la_primera_columna_y_es_unico_por_fila(self):
        f.make_inventory_item(self.db, name="Leche", code="0001")
        f.make_inventory_item(self.db, name="Fresa", code="0002")
        wb = self._export()
        codigos = [row[0].value for row in wb.active.iter_rows(min_row=2)]
        self.assertEqual(sorted(codigos), ["0001", "0002"])


class ExportLabelsTests(unittest.TestCase):
    """Historia 2 (P2): etiquetas de Tipo/Unidad/Estado idénticas a pantalla."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)
        self.admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")

    def _export(self):
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get("/api/v1/inventory/items/export")
        return openpyxl.load_workbook(io.BytesIO(resp.content))

    def test_tipo_raw_material_exporta_materia_prima_y_packaged_exporta_empacado(self):
        f.make_inventory_item(self.db, name="Leche", type="raw_material")
        f.make_inventory_item(self.db, name="Vaso desechable", type="packaged")
        wb = self._export()
        tipos = {row[1].value: row[2].value for row in wb.active.iter_rows(min_row=2)}
        self.assertEqual(tipos["Leche"], "Materia prima")
        self.assertEqual(tipos["Vaso desechable"], "Empacado")

    def test_unidad_exporta_la_abreviatura_no_el_nombre_largo(self):
        unit = f.make_unit(self.db, name="Litro", abbreviation="lt")
        f.make_inventory_item(self.db, name="Leche", unit=unit)
        wb = self._export()
        row = next(wb.active.iter_rows(min_row=2, max_row=2))
        self.assertEqual(row[3].value, "lt")

    def test_insumo_con_proveedor_preferido_exporta_solo_el_nombre(self):
        """spec 138, FR-001/FR-002: la fila de un insumo con proveedor preferido
        asignado trae el nombre del proveedor, pero ninguna columna de
        presentación/factor/último costo -- esos 3 datos ya no forman parte de
        la plantilla."""
        item = f.make_inventory_item(self.db, name="Con proveedor")
        supplier = f.make_supplier(self.db, name="Distribuidora Norte")
        f.make_supplier_item(
            self.db, supplier, item,
            presentation="Bolsa x 500 g", conversion_factor=Decimal("500"),
            last_cost=Decimal("4000"), preferred=True,
        )
        wb = self._export()
        row = [c.value for c in next(wb.active.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(len(row), 8)
        self.assertEqual(row[7], "Distribuidora Norte")

    def test_insumo_sin_proveedor_preferido_deja_esa_columna_en_blanco(self):
        """spec 109, Historia 1, escenario 3: sin proveedor preferido, la
        columna queda en blanco, sin error."""
        f.make_inventory_item(self.db, name="Sin proveedor")
        wb = self._export()
        row = next(wb.active.iter_rows(min_row=2, max_row=2))
        # openpyxl normaliza la cadena vacía a None al leerla de vuelta.
        self.assertFalse(row[7].value)


class ExportEncodingTests(unittest.TestCase):
    """Historia 2 (P2): nombres con tildes/eñes/comillas sin corrupción (FR-007)."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)
        self.admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")

    def test_nombres_con_caracteres_del_espanol_se_preservan_intactos(self):
        f.make_inventory_item(self.db, name="Champiñón")
        f.make_inventory_item(self.db, name="Arequipe (500ml)")
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get("/api/v1/inventory/items/export")

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        nombres = [row[1].value for row in wb.active.iter_rows(min_row=2)]
        self.assertIn("Champiñón", nombres)
        self.assertIn("Arequipe (500ml)", nombres)


def _xlsx_bytes(headers, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for fila in filas:
        ws.append(fila)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


class ImportEndpointTests(unittest.TestCase):
    """spec 109, Historia 2/3: POST /inventory/items/import -- estructura
    (FR-005), candado de una importación a la vez (FR-010) y encolado de la
    tarea Celery, sin necesitar un worker real (se mockea `.delay`, mismo
    patrón que `send_email_task` en `test_auth_forgot_password.py`)."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)
        self.admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")

    def _upload(self, content: bytes):
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        with mock.patch("app.api.v1.inventory.router.import_inventory_task") as mock_task:
            resp = client.post(
                "/api/v1/inventory/items/import",
                files={"file": ("inventario.xlsx", content,
                                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        return resp, mock_task

    def test_estructura_valida_crea_job_y_encola_la_tarea(self):
        content = _xlsx_bytes(INVENTORY_TEMPLATE_HEADERS, [
            ["", "Insumo nuevo", "Materia prima", "un", "0", "0", "0"],
        ])
        resp, mock_task = self._upload(content)
        self.assertEqual(resp.status_code, 202)
        body = resp.json()
        self.assertEqual(body["status"], "procesando")
        self.assertEqual(body["total_rows"], 1)
        mock_task.delay.assert_called_once()
        job = self.db.get(InventoryImportJob, UUID(body["id"]))
        self.assertIsNotNone(job)

    def test_archivo_legado_de_11_columnas_se_acepta_y_no_toca_presentacion_factor_ni_costo(self):
        """spec 138, Historia 1, escenario 2/FR-003/FR-004: un archivo exportado
        antes de este cambio (11 columnas, con presentación/factor/último costo
        al final) se sigue aceptando por estructura; esas 3 columnas se ignoran
        por completo -- aunque la fila cambie el proveedor preferido (columna
        que sí se conserva), la presentación/factor/último costo de la relación
        que queda preferida no cambian."""
        item = f.make_inventory_item(self.db, name="Con proveedor", code="0001")
        norte = f.make_supplier(self.db, name="Distribuidora Norte")
        otro = f.make_supplier(self.db, name="Otro Proveedor")
        f.make_supplier_item(
            self.db, norte, item,
            presentation="Bolsa x 500 g", conversion_factor=Decimal("500"),
            last_cost=Decimal("4000"), preferred=True,
        )
        relation_otro = f.make_supplier_item(
            self.db, otro, item,
            presentation="Caja x 10", conversion_factor=Decimal("10"),
            last_cost=None, preferred=False,
        )
        self.db.commit()

        legacy_headers = [
            *INVENTORY_TEMPLATE_HEADERS,
            "Presentación", "Factor de conversión", "Último costo",
        ]
        content = _xlsx_bytes(legacy_headers, [
            ["0001", "Con proveedor", "Materia prima", "un", "5", "1", "100",
             "Otro Proveedor", "Bulto x 25", "25", "9999"],
        ])
        resp, _mock_task = self._upload(content)
        self.assertEqual(resp.status_code, 202)
        body = resp.json()
        self.assertEqual(body["total_rows"], 1)

        job = self.db.get(InventoryImportJob, UUID(body["id"]))
        wb = openpyxl.load_workbook(io.BytesIO(content))
        filas = [[c.value for c in row] for row in wb.active.iter_rows(min_row=2)]
        import_service.procesar_archivo_importacion(self.db, job, filas)
        self.db.refresh(job)
        self.assertEqual(job.updated_count, 1, job.errors)
        self.assertEqual(job.failed_count, 0, job.errors)

        self.db.refresh(relation_otro)
        self.assertTrue(relation_otro.preferred, "debía marcarse como preferida")
        self.assertEqual(relation_otro.presentation, "Caja x 10")
        self.assertEqual(relation_otro.conversion_factor, Decimal("10"))
        self.assertIsNone(relation_otro.last_cost)

    def test_estructura_alterada_responde_422_sin_crear_job(self):
        content = _xlsx_bytes(["Nombre", "Código"], [["x", "y"]])
        resp, mock_task = self._upload(content)
        self.assertEqual(resp.status_code, 422)
        mock_task.delay.assert_not_called()
        count = self.db.execute(select(InventoryImportJob)).scalars().all()
        self.assertEqual(count, [])

    def test_archivo_vacio_responde_422(self):
        wb = openpyxl.Workbook()
        buffer = io.BytesIO()
        wb.save(buffer)
        resp, mock_task = self._upload(b"")
        self.assertEqual(resp.status_code, 422)
        mock_task.delay.assert_not_called()

    def test_segunda_importacion_con_una_en_curso_responde_409(self):
        self.db.add(InventoryImportJob(
            status="procesando", file_name="otro.xlsx",
            started_by_user_id=self.admin.id, total_rows=5,
        ))
        content = _xlsx_bytes(INVENTORY_TEMPLATE_HEADERS, [
            ["", "Insumo nuevo", "Materia prima", "un", "0", "0", "0"],
        ])
        resp, mock_task = self._upload(content)
        self.assertEqual(resp.status_code, 409)
        mock_task.delay.assert_not_called()

    def test_usuario_sin_rol_admin_no_puede_importar(self):
        cashier = pf.make_user(self.db, self.tenant, role_name="CASHIER")
        self.db.commit()
        client = _build_client(self.db, self.tenant, cashier)
        content = _xlsx_bytes(INVENTORY_TEMPLATE_HEADERS, [])
        with mock.patch("app.api.v1.inventory.router.import_inventory_task") as mock_task:
            resp = client.post(
                "/api/v1/inventory/items/import",
                files={"file": ("inventario.xlsx", content,
                                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        self.assertEqual(resp.status_code, 403)
        mock_task.delay.assert_not_called()

    def test_job_procesando_vencido_no_bloquea_y_queda_interrumpido(self):
        """spec 109, research.md D5: un job `procesando` de más de 30 minutos se
        considera abandonado -- no bloquea una importación nueva y queda marcado
        `interrumpido`."""
        viejo = InventoryImportJob(
            status="procesando", file_name="viejo.xlsx",
            started_by_user_id=self.admin.id, total_rows=1,
            started_at=utc_now() - timedelta(minutes=45),
        )
        self.db.add(viejo)
        self.db.commit()
        viejo_id = viejo.id

        content = _xlsx_bytes(INVENTORY_TEMPLATE_HEADERS, [
            ["", "Insumo nuevo", "Materia prima", "un", "0", "0", "0"],
        ])
        resp, mock_task = self._upload(content)
        self.assertEqual(resp.status_code, 202)
        mock_task.delay.assert_called_once()

        self.db.expire_all()
        self.assertEqual(self.db.get(InventoryImportJob, viejo_id).status, "interrumpido")

    def test_get_current_devuelve_el_job_mas_reciente(self):
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        self.assertIsNone(client.get("/api/v1/inventory/items/import/current").json())

        content = _xlsx_bytes(INVENTORY_TEMPLATE_HEADERS, [
            ["", "Insumo nuevo", "Materia prima", "un", "0", "0", "0"],
        ])
        resp, _mock_task = self._upload(content)
        job_id = resp.json()["id"]

        current = client.get("/api/v1/inventory/items/import/current").json()
        self.assertEqual(current["id"], job_id)
        self.assertEqual(current["status"], "procesando")
        self.assertFalse(current["has_errors"])

    def test_get_by_id_404_si_el_job_no_existe(self):
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get(f"/api/v1/inventory/items/import/{_uid()}")
        self.assertEqual(resp.status_code, 404)

    def test_descarga_de_errores_404_sin_filas_fallidas(self):
        job = InventoryImportJob(
            status="finalizado", file_name="ok.xlsx",
            started_by_user_id=self.admin.id, total_rows=1, failed_count=0,
        )
        self.db.add(job)
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get(f"/api/v1/inventory/items/import/{job.id}/errors")
        self.assertEqual(resp.status_code, 404)

    def test_descarga_de_errores_contiene_fila_codigo_y_motivo(self):
        job = InventoryImportJob(
            status="finalizado", file_name="con_errores.xlsx",
            started_by_user_id=self.admin.id, total_rows=1, failed_count=1,
            errors=[{"fila": 3, "codigo": "0099", "motivo": "Costo debe ser numérico"}],
        )
        self.db.add(job)
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get(f"/api/v1/inventory/items/import/{job.id}/errors")
        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        headers = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers, ["Fila", "Código", "Motivo"])
        row = [c.value for c in next(wb.active.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(row, [3, "0099", "Costo debe ser numérico"])


class ImportUnitLinkingTests(unittest.TestCase):
    """spec 138, Historia 2/3: la columna Unidad se vincula a una `UnitMeasure`
    real al crear un insumo nuevo (FR-006/FR-008/FR-009) y se ignora por completo
    al actualizar uno existente (FR-007). Llama `validar_y_aplicar_fila`
    directamente -- no pasa por el endpoint HTTP, igual que `fixtures.py`
    ejercita el resto de la lógica de negocio sin mocks."""

    def setUp(self):
        self.db = f.new_session()

    def _fila(self, code="", name="Insumo", tipo="Materia prima", unidad="un",
              stock="0", minimo="0", costo="0"):
        return [code, name, tipo, unidad, stock, minimo, costo]

    def test_unidad_en_minuscula_coincide_con_abreviatura_existente_en_mayuscula(self):
        """FR-006, Acceptance Scenario 1: "kg" vincula con la "KG" ya existente,
        sin crear una unidad duplicada."""
        f.make_unit(self.db, name="Kilogramo", abbreviation="KG")
        self.db.commit()
        antes = self.db.execute(select(UnitMeasure.id)).scalars().all()

        resultado = import_service.validar_y_aplicar_fila(
            self.db, self._fila(name="Harina", unidad="kg"),
        )
        self.assertTrue(resultado["ok"], resultado)
        self.assertEqual(resultado["accion"], "creado")

        despues = self.db.execute(select(UnitMeasure.id)).scalars().all()
        self.assertEqual(set(antes), set(despues), "no debía crear una unidad duplicada")
        item = self.db.execute(
            select(InventoryItem).where(InventoryItem.name == "Harina")
        ).scalar_one()
        unidad = self.db.get(UnitMeasure, item.unit_measure_id)
        self.assertEqual(unidad.abbreviation, "KG")

    def test_unidad_sin_coincidencia_crea_grupo_y_unidad_nueva_con_group_id(self):
        """FR-006, Acceptance Scenario 2 (research.md §3, el gap de `group_id`
        NOT NULL): "Galón" no existe -- se crea un `UnitMeasureGroup` propio y una
        `UnitMeasure` base de ese grupo, con `group_id` asignado (sin esto, la
        fila fallaría contra Postgres real por la restricción NOT NULL)."""
        resultado = import_service.validar_y_aplicar_fila(
            self.db, self._fila(name="Leche en galón", unidad="Galón"),
        )
        self.assertTrue(resultado["ok"], resultado)

        item = self.db.execute(
            select(InventoryItem).where(InventoryItem.name == "Leche en galón")
        ).scalar_one()
        unidad = self.db.get(UnitMeasure, item.unit_measure_id)
        self.assertEqual(unidad.name, "GALÓN")
        self.assertEqual(unidad.abbreviation, "GALÓN")
        self.assertTrue(unidad.is_base)
        self.assertEqual(unidad.factor_to_base, Decimal("1"))
        self.assertIsNotNone(unidad.group_id)
        grupo = self.db.get(UnitMeasureGroup, unidad.group_id)
        self.assertIsNotNone(grupo)
        self.assertEqual(grupo.name, "GALÓN")

    def test_dos_unidades_ambiguas_sin_distinguir_mayusculas_fallan_la_fila(self):
        """FR-008: "Gal" y "GAL" ya existen (dato ambiguo); una fila nueva con
        Unidad "gal" se reporta como fallida, en vez de elegir una al azar."""
        f.make_unit(self.db, name="Gal", abbreviation="Gal")
        f.make_unit(self.db, name="GAL-2", abbreviation="GAL")
        self.db.commit()

        resultado = import_service.validar_y_aplicar_fila(
            self.db, self._fila(name="Insumo ambiguo", unidad="gal"),
        )
        self.assertFalse(resultado["ok"])
        self.assertIn("más de una unidad", resultado["motivo"])
        existe = self.db.execute(
            select(InventoryItem).where(InventoryItem.name == "Insumo ambiguo")
        ).scalar_one_or_none()
        self.assertIsNone(existe)

    def test_unidad_vacia_en_fila_nueva_falla_por_campo_obligatorio(self):
        """FR-009: igual que cualquier otro campo requerido vacío."""
        resultado = import_service.validar_y_aplicar_fila(
            self.db, self._fila(name="Insumo sin unidad", unidad=""),
        )
        self.assertFalse(resultado["ok"])
        self.assertIn("obligatoria", resultado["motivo"])
        existe = self.db.execute(
            select(InventoryItem).where(InventoryItem.name == "Insumo sin unidad")
        ).scalar_one_or_none()
        self.assertIsNone(existe)


class ImportUnitUnchangedOnUpdateTests(unittest.TestCase):
    """spec 138, Historia 3 (FR-007): actualizar un insumo existente por
    importación nunca cambia su unidad base ni crea una `UnitMeasure`/
    `UnitMeasureGroup` nueva como efecto secundario, sin importar el valor de la
    columna Unidad de la fila."""

    def setUp(self):
        self.db = f.new_session()
        self.unit = f.make_unit(self.db, name="Gramo", abbreviation="g")
        self.item = f.make_inventory_item(
            self.db, name="Azúcar", code="0001", unit=self.unit,
        )
        self.db.commit()

    def _assert_unidad_sin_cambios_y_sin_unidades_nuevas(self, unidad_en_fila: str):
        antes = set(self.db.execute(select(UnitMeasure.id)).scalars().all())
        resultado = import_service.validar_y_aplicar_fila(
            self.db, [
                "0001", "Azúcar", "Materia prima", unidad_en_fila, "5", "1", "100",
            ],
        )
        self.assertTrue(resultado["ok"], resultado)
        self.assertEqual(resultado["accion"], "actualizado")

        self.db.refresh(self.item)
        self.assertEqual(self.item.unit_measure_id, self.unit.id)
        self.assertEqual(self.item.current_stock, Decimal("5"))
        self.assertEqual(self.item.unit_cost, Decimal("100"))

        despues = set(self.db.execute(select(UnitMeasure.id)).scalars().all())
        self.assertEqual(antes, despues, "no debía crear ninguna unidad nueva")

    def test_unidad_distinta_a_la_actual_se_ignora(self):
        f.make_unit(self.db, name="Kilogramo", abbreviation="KG")
        self.db.commit()
        self._assert_unidad_sin_cambios_y_sin_unidades_nuevas("KG")

    def test_unidad_vacia_se_ignora(self):
        self._assert_unidad_sin_cambios_y_sin_unidades_nuevas("")

    def test_unidad_inexistente_se_ignora_sin_crear_nada(self):
        self._assert_unidad_sin_cambios_y_sin_unidades_nuevas("Unidad que no existe")


class ImportPreferredSupplierTests(unittest.TestCase):
    """spec 138 (ajuste pedido por el usuario): la columna Proveedor preferido
    se conserva en export/import. El import vincula el proveedor (creándolo si
    no existe) y marca como preferida su relación con el insumo (creándola si no
    existe, con presentación/factor en `NULL`) -- nunca modifica presentación/
    factor de conversión/último costo de una relación ya existente."""

    def setUp(self):
        self.db = f.new_session()
        self.item = f.make_inventory_item(self.db, name="Insumo", code="0001")
        self.db.commit()

    def _fila(self, proveedor: str) -> list:
        return ["0001", "Insumo", "Materia prima", "un", "5", "1", "100", proveedor]

    def test_proveedor_con_relacion_existente_se_marca_preferida(self):
        norte = f.make_supplier(self.db, name="Distribuidora Norte")
        otro = f.make_supplier(self.db, name="Otro Proveedor")
        f.make_supplier_item(self.db, norte, self.item, preferred=True)
        relacion_otro = f.make_supplier_item(self.db, otro, self.item, preferred=False)
        self.db.commit()

        resultado = import_service.validar_y_aplicar_fila(
            self.db, self._fila("otro proveedor"),
        )
        self.assertTrue(resultado["ok"], resultado)

        self.db.refresh(relacion_otro)
        self.assertTrue(relacion_otro.preferred)

    def test_proveedor_sin_relacion_existente_se_crea_proveedor_y_relacion_preferida(self):
        """spec 138 (ajuste pedido por el usuario): el proveedor se crea si no
        existe, y la relación preferida también -- con presentación/factor en
        `NULL` porque la plantilla ya no los trae; el admin los completa luego
        desde el módulo de proveedores."""
        nombre = "Proveedor Nuevo"
        resultado = import_service.validar_y_aplicar_fila(
            self.db, self._fila(nombre),
        )
        self.assertTrue(resultado["ok"], resultado)

        proveedor = self.db.execute(
            select(Supplier).where(Supplier.name == nombre.upper())
        ).scalar_one()
        relacion = self.db.execute(
            select(SupplierItem).where(
                SupplierItem.inventory_item_id == self.item.id,
                SupplierItem.supplier_id == proveedor.id,
            )
        ).scalar_one()
        self.assertTrue(relacion.preferred)
        self.assertIsNone(relacion.presentation)
        self.assertIsNone(relacion.conversion_factor)

        self.db.refresh(self.item)
        self.assertEqual(self.item.unit_cost, Decimal("100"))

    def test_proveedor_vacio_no_toca_la_preferida_actual(self):
        norte = f.make_supplier(self.db, name="Distribuidora Norte")
        relacion = f.make_supplier_item(self.db, norte, self.item, preferred=True)
        self.db.commit()

        resultado = import_service.validar_y_aplicar_fila(self.db, self._fila(""))
        self.assertTrue(resultado["ok"], resultado)

        self.db.refresh(self.item)
        self.assertEqual(self.item.unit_cost, Decimal("100"))
        self.db.refresh(relacion)
        self.assertTrue(relacion.preferred)

    def test_dos_proveedores_con_el_mismo_nombre_sin_distinguir_mayusculas_fallan_la_fila(self):
        f.make_supplier(self.db, name="ABC")
        f.make_supplier(self.db, name="abc")
        self.db.commit()

        resultado = import_service.validar_y_aplicar_fila(self.db, self._fila("ABC"))
        self.assertFalse(resultado["ok"])
        self.assertIn("más de un proveedor", resultado["motivo"])

    def test_dos_relaciones_activas_con_el_mismo_proveedor_fallan_la_fila(self):
        """Ambigüedad al nivel de la relación (no del proveedor): el mismo
        proveedor tiene dos relaciones activas con este insumo (presentaciones
        distintas) -- no hay forma de elegir cuál marcar preferida."""
        proveedor = f.make_supplier(self.db, name="Distribuidora Norte")
        f.make_supplier_item(
            self.db, proveedor, self.item, presentation="Bolsa x 500 g", preferred=False,
        )
        f.make_supplier_item(
            self.db, proveedor, self.item, presentation="Caja x 10", preferred=False,
        )
        self.db.commit()

        resultado = import_service.validar_y_aplicar_fila(
            self.db, self._fila("Distribuidora Norte"),
        )
        self.assertFalse(resultado["ok"])
        self.assertIn("más de una relación", resultado["motivo"])


if __name__ == "__main__":
    unittest.main()
